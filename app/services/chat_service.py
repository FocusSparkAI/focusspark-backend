from pathlib import Path

from sqlmodel import Session

from app.ai.features.chat import chat_feature
from app.models.chat_model import ChatMessage, ChatThread, Document
from app.services.achievement_service import award_earned_achievements_for_user

ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}
ALLOWED_DOCUMENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "text/plain",
}


def _extract_text_from_pdf(file_bytes: bytes) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text.strip())
    return "\n\n".join(pages)


def _extract_text_from_docx(file_bytes: bytes) -> str:
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(file_bytes))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def _extract_text_from_pptx(file_bytes: bytes) -> str:
    from io import BytesIO

    from pptx import Presentation

    presentation = Presentation(BytesIO(file_bytes))
    slides = []
    for slide in presentation.slides:
        slide_parts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                slide_parts.append(text.strip())
        if slide_parts:
            slides.append("\n".join(slide_parts))
    return "\n\n".join(slides)


def _extract_text_from_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def extract_document_text(file_name: str, content_type: str | None, file_bytes: bytes) -> str:
    extension = Path(file_name).suffix.lower()
    if extension not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, PPTX, and TXT files are supported")

    if content_type and content_type not in ALLOWED_DOCUMENT_TYPES:
        raise ValueError("Only PDF, DOCX, PPTX, and TXT files are supported")

    if extension == ".pdf":
        text = _extract_text_from_pdf(file_bytes)
    elif extension == ".docx":
        text = _extract_text_from_docx(file_bytes)
    elif extension == ".pptx":
        text = _extract_text_from_pptx(file_bytes)
    else:
        text = _extract_text_from_txt(file_bytes)

    cleaned_text = text.strip()
    if not cleaned_text:
        raise ValueError("No readable text was found in the uploaded document")

    return cleaned_text


def generate_chat_response(
    message: str,
    provider_name: str | None = None,
    model_name: str | None = None,
):
    return chat_feature(message, provider_name=provider_name, model_name=model_name)


def generate_document_chat_response(message: str, document_text: str, document_name: str | None = None,
                                    provider_name: str | None = None, model_name: str | None = None):
    return chat_feature(message, document_text=document_text, document_name=document_name,
                        provider_name=provider_name, model_name=model_name)


def _resolve_model_for_provider(provider: str, model: str | None = None) -> str:
    if model and model.strip():
        return model.strip()

    normalized_provider = provider.strip().lower()
    from app.ai.settings import settings

    if normalized_provider == "gemini":
        return settings.gemini_model

    return settings.model


def create_thread(
    title: str | None,
    user_id: int,
    session: Session,
    ai_provider: str = "openai",
    ai_model: str | None = None,
):
    normalized_provider = (ai_provider or "openai").strip().lower()
    if normalized_provider not in {"openai", "gemini"}:
        normalized_provider = "openai"

    normalized_model = _resolve_model_for_provider(normalized_provider, ai_model)

    thread = ChatThread(
        user_id=user_id,
        title=title,
        ai_provider=normalized_provider,
        ai_model=normalized_model,
    )
    session.add(thread)
    session.commit()
    session.refresh(thread)
    return thread


def handle_chat(message: str, thread_id: int, user_id: int, session: Session):
    thread = session.get(ChatThread, thread_id)
    if not thread:
        raise ValueError("Thread not found")
    if thread.user_id != user_id:
        raise PermissionError("You do not have access to this thread")

    ai_response = generate_chat_response(
        message,
        provider_name=thread.ai_provider,
        model_name=thread.ai_model,
    )
    chat = ChatMessage(
        thread_id=thread_id,
        type="user",
        content=message
    )
    session.add(chat)
    session.commit()
    session.refresh(chat)
    ai_msg = ChatMessage(
        thread_id=thread_id,
        type="ai",
        content=ai_response
    )
    session.add(ai_msg)
    session.commit()
    session.refresh(ai_msg)
    return ai_msg


def handle_document_chat(
    message: str,
    thread_id: int,
    user_id: int,
    session: Session,
    file_name: str,
    content_type: str | None,
    file_bytes: bytes,
):
    thread = session.get(ChatThread, thread_id)
    if not thread:
        raise ValueError("Thread not found")
    if thread.user_id != user_id:
        raise PermissionError("You do not have access to this thread")

    document_text = extract_document_text(file_name, content_type, file_bytes)
    ai_response = generate_document_chat_response(message, document_text=document_text, document_name=file_name,
                                                 provider_name=thread.ai_provider, model_name=thread.ai_model)

    user_prompt = message.strip() or "Explain the uploaded document."
    user_message = ChatMessage(
        thread_id=thread_id,
        type="user",
        content=user_prompt,
        payload={
            "document_name": file_name,
            "document_type": content_type,
        },
    )
    session.add(user_message)
    session.commit()
    session.refresh(user_message)

    document = Document(
        user_id=user_id,
        name=file_name,
        file_type=content_type or Path(file_name).suffix.lower().lstrip(".") or "document",
        file_size=len(file_bytes),
        storage_path=f"chat-upload:{thread_id}:{user_message.id}:{file_name}",
        extracted_text=document_text,
        processed=True,
        uploaded_from="chat",
    )
    session.add(document)
    session.commit()
    award_earned_achievements_for_user(user_id, session)

    ai_msg = ChatMessage(
        thread_id=thread_id,
        type="ai",
        content=ai_response,
        payload={
            "document_name": file_name,
            "document_type": content_type,
        },
    )
    session.add(ai_msg)
    session.commit()
    session.refresh(ai_msg)
    return ai_msg
